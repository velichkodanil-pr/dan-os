"""R6.1A — emergency knowledge safety.

Every fixture in this file is a CLEARLY SYNTHETIC, NON-FUNCTIONAL credential:
made-up shapes that match a format but grant nothing anywhere. Nothing real
belongs in a test file, a log, a commit or an issue.

The provider tripwires below are the point of most of these tests. A gate that
"works" while still paying for an embedding call has not actually stopped the
leak — the secret already left the building. So the assertions are not only
about status fields: they fail the moment anything reaches out to OpenAI,
Anthropic or the network at all.
"""
import json

import httpx
import pytest
from sqlalchemy import func, select

from app.config import settings
from app.core import chat_tools, rag, secret_policy, security, security_scan, wiki
from app.core.ingest import ingest_document, ingest_document_parts
from app.core.orchestrator import Orchestrator
from app.core.secret_policy import SecretCategory, scan_text
from app.models import (
    ChatLog, Document, KnowledgeChunk, MemoryItem, Proposal, RawEvent,
    SecurityFinding, WikiPage,
)

OWNER = 111


def _shape(*parts: str) -> str:
    """Assemble a synthetic credential at run time.

    Every fixture in this file is fake, but a push-protection scanner reads
    source text and cannot know that — and it is right not to guess. So the
    provider-shaped literals never appear in the file at all: they are joined
    here. One `"".join` is a cheap price for never being tempted by a
    «Bypass» button on a repository whose whole point is not leaking tokens.
    """
    return "".join(parts)


# --- synthetic, non-functional fixtures (shape only, zero access) ---
FAKE_PASSWORD_LINE = "Логін: i.k@example.invalid | Пароль: Zx9-kLm2-Qw7"
FAKE_API_KEY = _shape("sk-", "ant-", "EXAMPLEONLY0000notarealkey1111AAAA")
FAKE_GITHUB_TOKEN = _shape("ghp", "_", "EXAMPLEONLYnotarealtoken00000000000")
FAKE_GOOGLE_KEY = _shape("AIza", "EXAMPLEONLYnotarealkey00000000000000")
FAKE_TELEGRAM_TOKEN = _shape("1234567890", ":", "AAEXAMPLEONLYnotarealtoken0000",
                             "00000")
FAKE_SLACK_TOKEN = _shape("xox", "b-", "1111111111-EXAMPLEONLYnotreal")
FAKE_OAUTH_SECRET = _shape("GOCSPX", "-", "EXAMPLEONLYnotreal00")
FAKE_JWT = _shape("eyJ", "hbGciOiJIUzI1NiJ9.", "eyJzdWIiOiJmYWtlIn0.",
                  "SIGNATUREfake00")
FAKE_PEM = ("-----BEGIN RSA PRIVATE KEY-----\n"
            "MIIEXAMPLEONLYnotarealkey0000\n-----END RSA PRIVATE KEY-----")
FAKE_BEARER = "Authorization: Bearer EXAMPLEONLYnotarealtoken0000"
# a line that still BLOCKS under the default policy (hard technical secret)
FAKE_SECRET_LINE = f"Сервіс api.example, ключ доступу {FAKE_API_KEY}"
BUSINESS_TEXT = ("Toco UA (ТОКО Україна) — партнер-оператор.\n\n"
                 "Сайт: toco-tour.example, менеджер i.k@example.invalid,\n\n"
                 "ЄДРПОУ 46140224, депозит 30%, комісія 12%.")


# ---------------------------------------------------------------- tripwires

class _DeadEmbedder:
    async def embed(self, texts):
        raise AssertionError("embedder must not be called on this path")


@pytest.fixture
def no_providers(monkeypatch):
    """Fail the test on ANY provider or network call."""
    async def _no_anthropic(*a, **kw):
        raise AssertionError("Anthropic must not be called on this path")

    class _NoNetwork:
        def __init__(self, *a, **kw):
            raise AssertionError("no network call is allowed on this path")

    monkeypatch.setattr("app.core.ingest.get_embedder", lambda: _DeadEmbedder())
    monkeypatch.setattr("app.core.rag.get_embedder", lambda: _DeadEmbedder())
    monkeypatch.setattr("app.core.extraction.haiku_text", _no_anthropic)
    monkeypatch.setattr(httpx, "AsyncClient", _NoNetwork)
    yield


class _TripwireExtractor:
    async def extract(self, text, context=None):
        raise AssertionError("extractor must not see this text")


@pytest.fixture(autouse=True)
def _reset_blocking():
    """Default policy: passwords allowed. Restore it after every test so a
    test that flips to strict mode cannot leak into the next one."""
    secret_policy.set_blocking_categories(secret_policy.HARD_SECRET_CATEGORIES)
    yield
    secret_policy.set_blocking_categories(secret_policy.HARD_SECRET_CATEGORIES)


# ============================================================ 1. the scanner

@pytest.mark.parametrize("text,category", [
    (FAKE_PEM, SecretCategory.PRIVATE_KEY),
    (FAKE_API_KEY, SecretCategory.API_KEY),
    (FAKE_GITHUB_TOKEN, SecretCategory.API_KEY),
    (FAKE_GOOGLE_KEY, SecretCategory.API_KEY),
    (FAKE_TELEGRAM_TOKEN, SecretCategory.API_KEY),
    (FAKE_SLACK_TOKEN, SecretCategory.API_KEY),
    (FAKE_BEARER, SecretCategory.BEARER_TOKEN),
    (FAKE_JWT, SecretCategory.BEARER_TOKEN),
    ('"refresh_token": "EXAMPLEONLYnotreal000"', SecretCategory.OAUTH_TOKEN),
    (f"client_secret={FAKE_OAUTH_SECRET}", SecretCategory.OAUTH_TOKEN),
    ("PHPSESSID=abcd1234efgh5678ijkl", SecretCategory.SESSION_COOKIE),
    ("Recovery codes: a1b2-c3d4, e5f6-g7h8, i9j0-k1l2, m3n4-o5p6",
     SecretCategory.RECOVERY_CODE),
    ("seed phrase: alpha bravo charlie delta echo foxtrot golf hotel india "
     "juliet kilo lima", SecretCategory.SEED_PHRASE),
])
def test_01_scanner_blocks_every_hard_secret_shape(text, category):
    result = scan_text(text)
    assert result.blocked, text[:40]
    assert category in result.categories


def test_01b_password_is_allowed_by_default():
    """Owner decision (R6.1A.1): a partner-portal password is searchable
    business data, not a blocked secret."""
    assert not scan_text(FAKE_PASSWORD_LINE).blocked
    assert not scan_text("Password: Qw3rty-Zx9-Lm").blocked
    assert not scan_text("Оператор | Пароль\nToco UA | Qw3rty-Zx9-Lm").blocked
    mixed = scan_text(FAKE_PASSWORD_LINE + "\n" + FAKE_API_KEY)
    assert mixed.blocked and mixed.categories == (SecretCategory.API_KEY,)


def test_01c_password_blocks_when_flag_on():
    """The stricter mode stays one env var away."""
    secret_policy.set_blocking_categories(secret_policy.ALL_CATEGORIES)
    assert scan_text(FAKE_PASSWORD_LINE).blocked


@pytest.mark.parametrize("text", [
    "яка політика паролів у компанії?",
    "Password policy: minimum 12 characters, rotate quarterly",
    "Пароль — не менше 12 символів, змінюємо щокварталу",
    "Логін: i.kornienko@example.invalid, сайт toco-tour.example",
    "IBAN UA213223130000026007233566001, ЄДРПОУ 46140224, ІПН 1234567890",
    "Рахунок №INV-2026-0041 на 2 346 EUR, заявка 59266",
    "Отримувач ТОВ ГЕПАРД, банк ПриватБанк, МФО 305299",
    "Контакт: Ірина, +380 67 123 45 67",
    "Тримай token і secret у менеджері паролів, а не в таблиці",
    BUSINESS_TEXT,
])
def test_02_scanner_leaves_business_content_searchable(text):
    assert not scan_text(text).blocked, text[:50]


@pytest.mark.parametrize("text", [
    "PASSWORD=<PASSWORD>", "API_KEY=${TOKEN}", "api_key: YOUR_API_KEY",
    "password: ***", "password: [REDACTED]", "PASSWORD=%TOKEN%",
    "пароль: змінено", "password: https://vault.example/item/42",
])
def test_03_scanner_ignores_placeholders_and_pointers(text):
    assert not scan_text(text).blocked, text


def test_04_scanner_never_exposes_the_value():
    """The API surface must be value-free — no excerpt, no hash, no encoding."""
    import hashlib
    secret = "Zx9-kLm2-Qw7"
    result = scan_text(FAKE_PASSWORD_LINE)
    blob = json.dumps({"repr": repr(result), "meta": result.as_meta()})
    assert secret not in blob
    assert hashlib.sha256(secret.encode()).hexdigest()[:12] not in blob
    assert set(result.as_meta()) == {"categories", "finding_count",
                                     "scanner_version"}


def test_05_scanner_reads_the_whole_document_not_a_prefix():
    """The credential block lives in the tail of the workbook — that was the
    original defect, so a prefix scan is not good enough."""
    filler = "\n\n".join(f"Рядок {i}: готель, 2 346 EUR, заявка 59266"
                         for i in range(4000))
    assert len(filler) > secret_policy.SECTION_CHARS * 3
    assert not scan_text(filler).blocked
    assert scan_text(filler + "\n\n" + FAKE_SECRET_LINE).blocked
    assert scan_text(FAKE_SECRET_LINE + "\n\n" + filler).blocked
    middle = filler[:len(filler) // 2] + FAKE_SECRET_LINE + filler[len(filler) // 2:]
    assert scan_text(middle).blocked


def test_06_scanner_normalises_unicode_evasion():
    secret_policy.set_blocking_categories(secret_policy.ALL_CATEGORIES)
    assert scan_text("Password: Qw3rty-Zx9-Lm").blocked      # nbsp
    assert scan_text("ｐａｓｓｗｏｒｄ：Qw3rty-Zx9-Lm").blocked        # fullwidth
    assert scan_text("pass​word: Qw3rty-Zx9-Lm").blocked     # zero width


def test_07_scanner_is_deterministic_local_code(monkeypatch):
    """No LLM, no embeddings, no web, no connectors — even if they exploded."""
    class _NoNetwork:
        def __init__(self, *a, **kw):
            raise AssertionError("the scanner must not touch the network")
    monkeypatch.setattr(httpx, "AsyncClient", _NoNetwork)
    first = scan_text(FAKE_SECRET_LINE + BUSINESS_TEXT)
    second = scan_text(FAKE_SECRET_LINE + BUSINESS_TEXT)
    assert first == second and first.blocked


def test_08_scan_result_is_frozen():
    result = scan_text(FAKE_API_KEY)
    with pytest.raises(Exception):
        result.blocked = False


# ============================================================ 2. ingest gate

@pytest.mark.asyncio
async def test_09_ingest_quarantines_without_any_provider_call(db, no_providers):
    result = await ingest_document(
        db, user_id=OWNER, domain="personal", title="Доступи DMC.xlsx",
        text=BUSINESS_TEXT + "\n\n" + FAKE_SECRET_LINE,
        source_type="telegram_file", source_ref="dmc.xlsx")
    assert result.status == "quarantined"
    assert result.chunks == 0
    assert result.document.status == "quarantined"


@pytest.mark.asyncio
async def test_10_quarantined_document_stores_no_source_text(db, no_providers):
    await ingest_document(db, user_id=OWNER, domain="personal", title="Доступи.xlsx",
                          text=BUSINESS_TEXT + "\n" + FAKE_SECRET_LINE,
                          source_type="drive", source_ref="f1")
    doc = (await db.execute(select(Document))).scalar_one()
    chunks = (await db.execute(
        select(func.count()).select_from(KnowledgeChunk))).scalar_one()
    assert chunks == 0 and doc.chunk_count == 0
    blob = json.dumps({"title": doc.title, "meta": doc.meta, "ref": doc.source_ref},
                      ensure_ascii=False)
    assert "EXAMPLEONLY" not in blob
    assert doc.meta["security"]["categories"] == ["api_key"]


@pytest.mark.asyncio
async def test_11_finding_is_metadata_only(db, no_providers):
    await ingest_document(db, user_id=OWNER, domain="personal", title="Доступи.txt",
                          text=FAKE_BEARER + "\n" + FAKE_API_KEY,
                          source_type="drive", source_ref="f2")
    finding = (await db.execute(select(SecurityFinding))).scalar_one()
    assert finding.status == "open"
    assert set(finding.categories) == {"bearer_token", "api_key"}
    assert finding.finding_count >= 2
    columns = {c.name for c in SecurityFinding.__table__.columns}
    assert not columns & {"excerpt", "value", "secret", "hash", "fingerprint",
                          "text", "payload"}
    assert "EXAMPLEONLY" not in json.dumps(
        {c: str(getattr(finding, c)) for c in columns}, ensure_ascii=False)


@pytest.mark.asyncio
async def test_12_repeat_ingest_is_idempotent(db, no_providers):
    text = BUSINESS_TEXT + "\n" + FAKE_SECRET_LINE
    first = await ingest_document(db, user_id=OWNER, domain="personal",
                                  title="a.txt", text=text,
                                  source_type="drive", source_ref="f3")
    second = await ingest_document(db, user_id=OWNER, domain="personal",
                                   title="a.txt", text=text,
                                   source_type="drive", source_ref="f3")
    assert first.status == second.status == "quarantined"
    assert (await db.execute(
        select(func.count()).select_from(Document))).scalar_one() == 1
    assert (await db.execute(
        select(func.count()).select_from(SecurityFinding))).scalar_one() == 1


@pytest.mark.asyncio
async def test_13_ordinary_document_still_indexes(db):
    """The round must not turn DAN.OS into a redaction machine."""
    result = await ingest_document(
        db, user_id=OWNER, domain="personal", title="Умови операторів",
        text=BUSINESS_TEXT + "\n\n" + "\n\n".join(
            f"Оператор {i}: депозит {i}%, ЄДРПОУ 4614022{i}" for i in range(20)),
        source_type="drive", source_ref="ok1")
    assert result.status == "indexed" and result.chunks > 0
    assert result.document.status == "indexed"
    assert not (await db.execute(select(SecurityFinding))).scalars().all()


@pytest.mark.asyncio
async def test_14_parts_gate_contains_the_whole_source(db, no_providers):
    """A secret must not be smuggled in by landing on a part boundary."""
    from app.core.ingest import PART_CHARS
    body = "\n\n".join(f"Рядок даних номер {i} з довгим описом усередині"
                       for i in range(9000))
    assert len(body) > PART_CHARS
    results = await ingest_document_parts(
        db, user_id=OWNER, domain="personal", title="Великий.xlsx",
        text=body + "\n\n" + FAKE_SECRET_LINE,
        source_type="drive", source_ref="big1")
    assert [r.status for r in results] == ["quarantined"]
    assert (await db.execute(
        select(func.count()).select_from(KnowledgeChunk))).scalar_one() == 0


@pytest.mark.asyncio
async def test_15_xlsx_quarantines_only_the_affected_sheet(db):
    import io
    from openpyxl import Workbook
    from app.core.ingest import ingest_xlsx_by_sheets
    wb = Workbook()
    ws = wb.active
    ws.title = "Продукт"
    for i in range(40):
        ws.append([f"Ідея {i}", f"опис ідеї номер {i} для сезону"])
    ws2 = wb.create_sheet("DMC")
    ws2.append(["Сервіс", "Ключ API"])
    ws2.append(["Travelon AI", FAKE_API_KEY])
    buf = io.BytesIO()
    wb.save(buf)
    results = await ingest_xlsx_by_sheets(
        db, user_id=OWNER, domain="personal", filename="Travelon.xlsx",
        data=buf.getvalue(), source_type="drive", source_ref="TP1")
    by_status = {r.document.title.split("«")[1][:4]: r.status
                 for r in results if r.document}
    assert by_status["Прод"] == "indexed"
    assert by_status["DMC»"] == "quarantined"


# ============================================================ 3. retrieval

@pytest.mark.asyncio
async def test_16_retrieval_excludes_quarantined_documents(db):
    result = await ingest_document(
        db, user_id=OWNER, domain="personal", title="Умови", text=BUSINESS_TEXT,
        source_type="drive", source_ref="q1")
    assert await rag.retrieve(db, user_id=OWNER, domain="personal",
                              query="ТОКО депозит комісія")
    result.document.status = "quarantined"
    await db.commit()
    assert await rag.retrieve(db, user_id=OWNER, domain="personal",
                              query="ТОКО депозит комісія") == []


@pytest.mark.asyncio
async def test_17_retrieval_withholds_a_legacy_secret_chunk(db):
    """Chunks indexed before R6.1A exist until the owner runs the scan —
    retrieval has to be safe on the first request after deploy, not later."""
    from app.core.embeddings import get_embedder
    doc = Document(user_id=OWNER, domain="personal", title="Стара таблиця",
                   source_type="drive", source_ref="legacy",
                   content_hash="legacyhash", status="indexed", chunk_count=1)
    db.add(doc)
    await db.flush()
    text = f"Toco UA депозит комісія {FAKE_SECRET_LINE}"
    emb = (await get_embedder().embed([text]))[0]
    db.add(KnowledgeChunk(document_id=doc.id, user_id=OWNER, seq=0,
                          text=text, embedding=emb))
    await db.commit()
    hits = await rag.retrieve(db, user_id=OWNER, domain="personal", query=text)
    assert hits == []


@pytest.mark.asyncio
async def test_18_keyword_fallback_survives_for_identifiers(db):
    await ingest_document(
        db, user_id=OWNER, domain="personal", title="Реквізити",
        text="Other | Toco UA | toco-tour.example | ЄДРПОУ 46140224\n\n"
             + "\n\n".join(f"Нейтральний рядок {i}" for i in range(30)),
        source_type="drive", source_ref="r1")
    hits = await rag.retrieve(db, user_id=OWNER, domain="personal",
                              query="реквізити ТОКО")
    assert any("46140224" in h.text for h in hits)


def test_19_credential_words_are_no_longer_lookup_triggers():
    for query in ("який пароль до Toco", "login and password", "логін до ТОКО"):
        assert not rag._LOOKUP_RE.search(query), query
    for query in ("реквізити ТОКО", "IBAN партнера", "ЄДРПОУ Гепард"):
        assert rag._LOOKUP_RE.search(query), query


# ============================================================ 4. chat intake

@pytest.mark.asyncio
async def test_20_note_with_secret_never_becomes_a_row(db, no_providers):
    orch = Orchestrator(extractor=_TripwireExtractor())
    outcome = await orch.handle_note(
        db, user_id=OWNER, text=f"збережи: {FAKE_SECRET_LINE}",
        dedupe_key="tg:1:1")
    assert outcome.kind == "blocked"
    for model in (RawEvent, ChatLog, MemoryItem, Proposal):
        assert (await db.execute(
            select(func.count()).select_from(model))).scalar_one() == 0, model
    assert (await db.execute(
        select(func.count()).select_from(SecurityFinding))).scalar_one() == 1


@pytest.mark.asyncio
async def test_21_blocked_reply_echoes_nothing(db, no_providers):
    orch = Orchestrator(extractor=_TripwireExtractor())
    outcome = await orch.handle_note(
        db, user_id=OWNER, text=f"ось ключ {FAKE_API_KEY}", dedupe_key="tg:1:2")
    assert "EXAMPLEONLY" not in (outcome.reply or "")
    assert "менеджер" in (outcome.reply or "").lower()
    from app.models import AuditRecord
    audit_rows = (await db.execute(select(AuditRecord))).scalars().all()
    blob = json.dumps([r.details for r in audit_rows], ensure_ascii=False)
    assert "EXAMPLEONLY" not in blob
    assert any(r.outcome == "denied" for r in audit_rows)


@pytest.mark.asyncio
async def test_22_password_policy_question_still_works(db, monkeypatch):
    """«яка політика паролів?» is an ordinary knowledge question."""
    assert not security.is_credential_request("яка політика паролів у нас?")
    assert not scan_text("яка політика паролів у нас?").blocked

    class _Reply:
        async def extract(self, text, context=None):
            from app.core.extraction import ExtractResult
            return ExtractResult(intent="chat", reply="Політика: 12+ символів.")
    monkeypatch.setattr("app.core.chat.chat_reply",
                        lambda *a, **kw: _async_none())
    orch = Orchestrator(extractor=_Reply())
    outcome = await orch.handle_note(db, user_id=OWNER,
                                     text="яка політика паролів у нас?",
                                     dedupe_key="tg:1:3")
    assert outcome.kind == "chat"
    assert (await db.execute(
        select(func.count()).select_from(RawEvent))).scalar_one() == 1


async def _async_none():
    return None


@pytest.mark.asyncio
async def test_23_token_request_refused_password_request_flows(db, monkeypatch):
    """A HARD-secret lookup is still refused without a model call. A password
    lookup is now an ordinary question — it must reach normal retrieval/chat."""
    orch = Orchestrator(extractor=_TripwireExtractor())
    monkeypatch.setattr("app.core.extraction.haiku_text",
                        lambda *a, **kw: _raise())
    out_token = await orch.handle_note(db, user_id=OWNER,
                                       text="дай токен доступу до бота",
                                       dedupe_key="tg:1:4a")
    assert out_token.kind == "chat"
    assert out_token.reply == security.SAFE_NOT_STORED

    assert not security.is_credential_request("який пароль до ТОКО Україна?")

    class _Reply:
        async def extract(self, text, context=None):
            from app.core.extraction import ExtractResult
            return ExtractResult(intent="chat", reply="Дивлюсь у базі…")
    monkeypatch.setattr("app.core.chat.chat_reply",
                        lambda *a, **kw: _async_none())
    orch2 = Orchestrator(extractor=_Reply())
    out_pw = await orch2.handle_note(db, user_id=OWNER,
                                     text="який пароль до ТОКО Україна?",
                                     dedupe_key="tg:1:4b")
    assert out_pw.kind == "chat"
    assert out_pw.reply != security.SAFE_NOT_STORED


async def _raise():
    raise AssertionError("no model call on a refused hard-secret lookup")


# ============================================================ 5. tools & wiki

@pytest.mark.asyncio
async def test_24_wiki_save_answer_is_gone(db):
    names = {t["name"] for t in chat_tools.TOOL_DEFS}
    assert "wiki_save_answer" not in names
    assert "wiki_save_answer" not in chat_tools._POLICY
    assert "wiki_save_answer" not in chat_tools._EXECUTORS
    denied = json.loads(await chat_tools.run_tool(
        db, OWNER, "personal", "wiki_save_answer",
        {"title": "x", "summary": "y", "body": "z"}))
    assert "не дозволено" in denied["error"]
    from app.core.chat import _SYSTEM
    assert "wiki_save_answer" not in _SYSTEM


@pytest.mark.asyncio
async def test_25_tool_output_is_withheld_when_it_carries_a_secret(db, monkeypatch):
    async def leaky(_db, _user_id, _domain, _args):
        return {"open_tasks": [{"title": f"ключ {FAKE_API_KEY}"}]}
    monkeypatch.setitem(chat_tools._EXECUTORS, "get_tasks", leaky)
    raw = await chat_tools.run_tool(db, OWNER, "personal", "get_tasks", {})
    assert "EXAMPLEONLY" not in raw
    assert json.loads(raw)["withheld"] is True
    finding = (await db.execute(select(SecurityFinding))).scalar_one()
    assert finding.resource_type == "tool_output"


@pytest.mark.asyncio
async def test_26_quarantined_page_is_invisible_to_every_reader(db):
    page, _ = await wiki.upsert_page(
        db, user_id=OWNER, domain="personal", kind="entity", title="Toco UA",
        summary="оператор", content="- Депозит 30%", aliases=["ТОКО"],
        tags=["partner"])
    await db.commit()
    assert await wiki.find_page(db, OWNER, "personal", "ТОКО") is not None
    page.status = "quarantined"
    await db.commit()
    assert await wiki.find_page(db, OWNER, "personal", "ТОКО") is None
    assert await wiki.search_pages(db, OWNER, "personal", "Toco") == []
    assert "Toco UA" not in await wiki.render_index(db, OWNER, "personal")
    assert (await wiki.lint(db, OWNER, "personal"))["quarantined"] == 1


@pytest.mark.asyncio
async def test_27_compiler_never_sends_a_secret_source(db, no_providers):
    outcome = await wiki.compile_source(
        db, user_id=OWNER, domain="personal", title="Доступи DMC",
        text=BUSINESS_TEXT + "\n" + FAKE_SECRET_LINE, source_ref="d1")
    assert outcome.status == "quarantined"
    assert outcome.pages == [] and outcome.error_code == "secret_detected"
    assert (await db.execute(
        select(func.count()).select_from(WikiPage))).scalar_one() == 0


@pytest.mark.asyncio
async def test_28_compiler_drops_secret_facts_from_model_output(db, monkeypatch):
    """The prompt forbids secret values; this is the enforcement behind it."""
    async def sloppy(prompt, max_tokens=600):
        return json.dumps({"pages": [{
            "kind": "entity", "title": "Toco UA", "aliases": ["ТОКО"],
            "summary": "оператор", "tags": ["partner"],
            "facts": ["Сайт: toco-tour.example", "Депозит: 30%",
                      FAKE_PASSWORD_LINE, f"Ключ: {FAKE_API_KEY}"]}]},
            ensure_ascii=False)
    monkeypatch.setattr("app.core.extraction.haiku_text", sloppy)
    outcome = await wiki.compile_source(db, user_id=OWNER, domain="personal",
                                        title="Партнери",
                                        text=BUSINESS_TEXT, source_ref="d2")
    assert outcome.status == "succeeded"
    page = await wiki.find_page(db, OWNER, "personal", "ТОКО")
    assert page.status == "active"
    assert "Депозит: 30%" in page.content
    assert "EXAMPLEONLY" not in page.content       # the api-key fact is dropped
    assert "Zx9-kLm2-Qw7" in page.content          # the password fact is kept


@pytest.mark.asyncio
async def test_29_compile_status_is_structured_and_queue_is_honest(db, monkeypatch):
    doc = await _indexed_doc(db, "Звичайний документ", BUSINESS_TEXT * 3)

    async def broken(prompt, max_tokens=600):
        return None
    monkeypatch.setattr("app.core.extraction.haiku_text", broken)
    outcome = await wiki.compile_document(db, user_id=OWNER, document=doc,
                                          domain="personal")
    assert outcome.status == "failed" and outcome.error_code == "provider_unavailable"
    state = wiki.compile_state(doc)
    assert state["compiler_version"] == wiki.COMPILER_VERSION
    assert state["source_chars"] > 0 and "at" in state
    # a provider blip must not remove a source from the base forever
    assert doc.id in {d.id for d in await wiki.pending_documents(db, OWNER, "personal")}

    async def empty(prompt, max_tokens=600):
        return json.dumps({"pages": []})
    monkeypatch.setattr("app.core.extraction.haiku_text", empty)
    outcome = await wiki.compile_document(db, user_id=OWNER, document=doc,
                                          domain="personal")
    assert outcome.status == "empty_valid"
    assert doc.id not in {d.id for d in await wiki.pending_documents(db, OWNER, "personal")}


@pytest.mark.asyncio
async def test_30_oversized_source_reports_deferred_large(db, monkeypatch):
    """«Compiled» must not quietly mean «read the first 12k characters»."""
    big = "\n\n".join(f"Оператор {i}: депозит {i}%, комісія {i}%"
                      for i in range(3000))
    assert len(big) > wiki.MAX_SOURCE_CHARS
    doc = await _indexed_doc(db, "Величезний реєстр", big)

    async def ok(prompt, max_tokens=600):
        return json.dumps({"pages": [{
            "kind": "entity", "title": "Реєстр операторів", "aliases": [],
            "summary": "s", "facts": ["Депозит: 30%"], "tags": []}]},
            ensure_ascii=False)
    monkeypatch.setattr("app.core.extraction.haiku_text", ok)
    outcome = await wiki.compile_document(db, user_id=OWNER, document=doc,
                                          domain="personal")
    assert outcome.status == "deferred_large"
    assert outcome.processed_chars < outcome.source_chars
    assert doc.id in {d.id for d in await wiki.pending_documents(db, OWNER, "personal")}


@pytest.mark.asyncio
async def test_31_upsert_contains_a_secret_page_on_write(db):
    page, status = await wiki.upsert_page(
        db, user_id=OWNER, domain="personal", kind="entity", title="Партнер X",
        summary="оператор", content=f"- {FAKE_SECRET_LINE}", aliases=[], tags=[])
    await db.commit()
    assert status == "quarantined" and page.status == "quarantined"
    assert page.content == "" and page.summary == ""
    finding = (await db.execute(select(SecurityFinding))).scalar_one()
    assert finding.resource_type == "wiki_page"


# ============================================================ 6. the DB scan

async def _indexed_doc(db, title, text, *, ref="doc"):
    result = await ingest_document(db, user_id=OWNER, domain="personal",
                                   title=title, text=text,
                                   source_type="drive", source_ref=ref)
    return result.document


@pytest.mark.asyncio
async def test_32_scan_contains_existing_content_and_is_idempotent(db, monkeypatch):
    """Legacy rows written by the pre-R6.1A pipeline get contained, not deleted."""
    doc = Document(user_id=OWNER, domain="personal", title="Стара таблиця",
                   source_type="drive", source_ref="legacy",
                   content_hash="h-legacy", status="indexed", chunk_count=1)
    db.add(doc)
    await db.flush()
    from app.core.embeddings import get_embedder
    emb = (await get_embedder().embed([FAKE_API_KEY]))[0]
    db.add(KnowledgeChunk(document_id=doc.id, user_id=OWNER, seq=0,
                          text=FAKE_API_KEY, embedding=emb))
    page = WikiPage(user_id=OWNER, kind="entity", slug="legacy", title="Legacy",
                    summary="s", content=FAKE_BEARER, contradictions="",
                    aliases=[], tags=[], sources=[], status="active")
    db.add(page)
    db.add(MemoryItem(user_id=OWNER, content=f"ключ {FAKE_API_KEY}",
                      status="confirmed"))
    db.add(ChatLog(user_id=OWNER, role="user", text=FAKE_API_KEY))
    db.add(RawEvent(event_type="telegram.message", dedupe_key="legacy-1",
                    user_id=OWNER, payload={"text": FAKE_API_KEY}))
    await db.commit()

    # a provider tripwire covering the whole scan
    async def _boom(*a, **kw):
        raise AssertionError("the scan must be fully local")
    monkeypatch.setattr("app.core.extraction.haiku_text", _boom)
    monkeypatch.setattr("app.core.rag.get_embedder", lambda: _DeadEmbedder())
    monkeypatch.setattr(httpx, "AsyncClient",
                        lambda *a, **kw: (_ for _ in ()).throw(
                            AssertionError("no network in the scan")))

    report = await security_scan.run_scan(db, user_id=OWNER)
    assert report.completed
    assert report.documents_quarantined == 1
    assert report.wiki_quarantined == 1
    assert report.memory_quarantined == 1
    assert report.chat_contained == 1
    assert report.raw_events_flagged == 1

    await db.refresh(doc); await db.refresh(page)
    assert doc.status == "quarantined" and page.status == "quarantined"
    # containment, NOT deletion — every row is still there
    assert (await db.execute(
        select(func.count()).select_from(KnowledgeChunk))).scalar_one() == 1
    assert (await db.execute(
        select(func.count()).select_from(RawEvent))).scalar_one() == 1
    event = (await db.execute(select(RawEvent))).scalar_one()
    assert event.payload["text"] == FAKE_API_KEY  # immutable, untouched

    findings_first = (await db.execute(
        select(func.count()).select_from(SecurityFinding))).scalar_one()
    second = await security_scan.run_scan(db, user_id=OWNER)
    assert second.completed
    assert (await db.execute(
        select(func.count()).select_from(SecurityFinding))).scalar_one() \
        == findings_first
    assert second.documents_quarantined == 0  # already contained


@pytest.mark.asyncio
async def test_33_interrupted_scan_leaves_the_gate_closed(db, monkeypatch):
    await _indexed_doc(db, "Умови", BUSINESS_TEXT, ref="ok")
    await security.mark_scan_complete(db)
    await db.commit()
    assert await security.scan_complete(db)

    async def explode(*a, **kw):
        raise RuntimeError("interrupted")
    monkeypatch.setattr(security_scan, "_scan_wiki", explode)
    with pytest.raises(RuntimeError):
        await security_scan.run_scan(db, user_id=OWNER)
    await db.rollback()
    assert not await security.scan_complete(db)


@pytest.mark.asyncio
async def test_34_scan_report_carries_counts_only(db):
    # a legacy row, written the way the pre-R6.1A pipeline wrote them
    from app.core.embeddings import get_embedder
    doc = Document(user_id=OWNER, domain="personal",
                   title="Дуже Секретний Файл.xlsx", source_type="drive",
                   source_ref="s1", content_hash="h-s1", status="indexed",
                   chunk_count=1)
    db.add(doc)
    await db.flush()
    emb = (await get_embedder().embed([FAKE_API_KEY]))[0]
    db.add(KnowledgeChunk(document_id=doc.id, user_id=OWNER, seq=0,
                          text=FAKE_API_KEY, embedding=emb))
    await db.commit()
    report = await security_scan.run_scan(db, user_id=OWNER)
    assert report.documents_quarantined == 1
    text = security_scan.report_text(report)
    assert "EXAMPLEONLY" not in text
    assert "Дуже Секретний Файл" not in text
    assert "перевипустити" in text.lower()


@pytest.mark.asyncio
async def test_35_auto_compilation_is_off_and_gated(db):
    assert settings.auto_wiki_compile_enabled is False
    assert not await security.scan_complete(db)          # closed until scanned
    await security_scan.run_scan(db, user_id=OWNER)
    assert await security.scan_complete(db)
    from app.main import AdminIngestRequest
    assert AdminIngestRequest(title="t", text="x" * 30).compile is False


# =========================================== 7. the other provider doorways

@pytest.mark.asyncio
async def test_36_gmail_digest_drops_a_credential_mail(monkeypatch):
    """A mailbox is external content; the digest is a provider call over it."""
    from app.core import digest
    monkeypatch.setattr(settings, "anthropic_api_key", "sk-test-not-real")

    class _NoNetwork:
        def __init__(self, *a, **kw):
            raise AssertionError("nothing may be sent for this mailbox")
    monkeypatch.setattr("app.core.digest.httpx.AsyncClient", _NoNetwork)

    only_secret = [{"from": "noreply@example.invalid", "subject": "Ваш ключ",
                    "snippet": FAKE_API_KEY}]
    assert digest._digestible(only_secret) == []
    assert await digest._rank_with_haiku(only_secret) is None

    mixed = only_secret + [{"from": "ira@example.invalid",
                            "subject": "Умови на сезон", "snippet": "депозит 30%"}]
    kept = digest._digestible(mixed)
    assert len(kept) == 1 and kept[0]["subject"] == "Умови на сезон"


@pytest.mark.asyncio
async def test_37_reply_draft_refuses_a_credential_letter(db, monkeypatch):
    from app.core import google_client
    from app.models import GoogleCredential, PendingDraft
    db.add(GoogleCredential(user_id=OWNER, account_email="me@example.invalid",
                            label="me", domain="personal", refresh_token_enc="enc"))
    await db.commit()

    async def _access(_db, _cred):
        return "token"

    async def _find(_access, _query):
        return {"from": "noreply@example.invalid", "subject": "Ваш новий ключ",
                "body": FAKE_API_KEY, "thread_id": "t", "message_id": "m",
                "references": ""}

    async def _no_model(*a, **kw):
        raise AssertionError("the letter must not reach the model")
    monkeypatch.setattr(google_client, "access_for", _access)
    monkeypatch.setattr(google_client, "gmail_find_message", _find)
    monkeypatch.setattr("app.core.extraction.haiku_text", _no_model)

    orch = Orchestrator(extractor=_TripwireExtractor())
    status, draft = await orch.propose_draft(db, user_id=OWNER, query="доступ")
    assert status == "blocked_secret" and draft is None
    assert (await db.execute(
        select(func.count()).select_from(PendingDraft))).scalar_one() == 0


@pytest.mark.asyncio
async def test_38_assembled_context_is_checked_as_a_whole(db, monkeypatch):
    """Chunks, calendar titles and profile facts are filtered upstream; the
    block that actually reaches the model is checked once more."""
    seen = {}

    class _Capture:
        async def extract(self, text, context=None):
            from app.core.extraction import ExtractResult
            seen.update(context or {})
            return ExtractResult(intent="chat", reply="ок")

    db.add(MemoryItem(user_id=OWNER, content=f"мій ключ {FAKE_API_KEY}",
                      status="confirmed"))
    db.add(MemoryItem(user_id=OWNER, content="Данило керує TravelON",
                      status="confirmed"))
    await db.commit()

    async def _agenda(_db, _user_id, days=7):
        return f"\nКалендар: зустріч «{FAKE_API_KEY}»"
    monkeypatch.setattr("app.core.briefs.agenda_block", _agenda)
    monkeypatch.setattr("app.core.chat.chat_reply", lambda *a, **kw: _async_none())

    orch = Orchestrator(extractor=_Capture())
    await orch.handle_note(db, user_id=OWNER, text="що в календарі на завтра?",
                           dedupe_key="tg:2:1")
    assert "EXAMPLEONLY" not in json.dumps(seen, ensure_ascii=False, default=str)
    assert "Zx9-kLm2-Qw7" not in json.dumps(seen, ensure_ascii=False, default=str)
    assert "Данило керує TravelON" in seen["profile"]


# ============================================ 8. /kb_quarantine listing

@pytest.mark.asyncio
async def test_39_quarantine_listing_names_sources_without_content(db):
    """The rotation walk-list: titles/dates/categories, never content."""
    # one file in three parts + one clean file + a quarantined wiki page
    for i in (1, 2, 3):
        db.add(Document(user_id=OWNER, domain="personal",
                        title=f"Доступи DMC (ч.{i})", source_type="drive",
                        source_ref="dmc", content_hash=f"h-dmc-{i}",
                        status="quarantined", chunk_count=0,
                        meta={"security": {"categories": ["password"],
                                           "finding_count": 2,
                                           "scanner_version": 1}}))
    db.add(Document(user_id=OWNER, domain="personal", title="Умови операторів",
                    source_type="drive", source_ref="ok", content_hash="h-ok",
                    status="indexed", chunk_count=3))
    db.add(WikiPage(user_id=OWNER, kind="entity", slug="toco", title="Toco UA",
                    summary="", content="", contradictions="", aliases=[],
                    tags=[], sources=[], status="quarantined"))
    db.add(ChatLog(user_id=OWNER, role="bot", text="стара репліка",
                   provider_eligible=False))
    await db.commit()

    listing = await security_scan.quarantine_listing(db, OWNER)
    assert listing["doc_rows"] == 3
    assert len(listing["documents"]) == 1          # parts merged into one line
    entry = listing["documents"][0]
    assert entry["title"] == "Доступи DMC" and entry["parts"] == 3
    assert entry["categories"] == ["password"]
    assert [p["title"] for p in listing["wiki"]] == ["Toco UA"]
    assert listing["chat_contained"] == 1

    messages = security_scan.quarantine_text(listing)
    joined = "\n".join(messages)
    assert "Доступи DMC" in joined and "Toco UA" in joined
    assert "Умови операторів" not in joined        # active docs stay out
    assert "стара репліка" not in joined           # chat text never shown
    assert all(len(m) <= 3500 for m in messages)


@pytest.mark.asyncio
async def test_40_quarantine_listing_masks_secret_bearing_titles(db):
    """A filename can itself carry the secret — the listing must not echo it."""
    db.add(Document(user_id=OWNER, domain="personal",
                    title=f"Нотатка {FAKE_API_KEY}", source_type="telegram_file",
                    source_ref="n1", content_hash="h-n1",
                    status="quarantined", chunk_count=0))
    await db.commit()
    listing = await security_scan.quarantine_listing(db, OWNER)
    joined = "\n".join(security_scan.quarantine_text(listing))
    assert "EXAMPLEONLY" not in joined
    assert "назву приховано" in joined


@pytest.mark.asyncio
async def test_41_quarantine_listing_empty_is_honest(db):
    listing = await security_scan.quarantine_listing(db, OWNER)
    messages = security_scan.quarantine_text(listing)
    assert len(messages) == 1 and "порожній" in messages[0]


# ==================================== 9. reconcile: release passwords, keep tokens

@pytest.mark.asyncio
async def test_42_rescan_releases_password_content_keeps_tokens(db):
    """After the owner allowed passwords, re-running the scan must RELEASE the
    password-only content it had quarantined, while a token document stays put.
    Documents keep their chunks from the pre-R6.1A indexing, so releasing them
    restores searchability."""
    from app.core.embeddings import get_embedder

    async def _quarantined_doc(title, ref, text, chash):
        doc = Document(user_id=OWNER, domain="personal", title=title,
                       source_type="drive", source_ref=ref, content_hash=chash,
                       status="quarantined", chunk_count=1,
                       meta={"security": {"categories": ["x"], "finding_count": 1,
                                          "scanner_version": 1}})
        db.add(doc)
        await db.flush()
        emb = (await get_embedder().embed([text]))[0]
        db.add(KnowledgeChunk(document_id=doc.id, user_id=OWNER, seq=0,
                              text=text, embedding=emb))
        await security.record_finding(
            db, user_id=OWNER, resource_type="document", resource_id=doc.id,
            result=security.SecretScanResult(True, (security.SecretCategory.PASSWORD,), 1))
        return doc

    pw_doc = await _quarantined_doc(
        "Партнери пароль", "pw", f"Toco UA депозит {FAKE_PASSWORD_LINE}", "h-pw")
    tok_doc = await _quarantined_doc(
        "Токени", "tok", f"Сервіс {FAKE_API_KEY}", "h-tok")
    # a quarantined password wiki page + a contained password chat line
    page = WikiPage(user_id=OWNER, kind="entity", slug="toco", title="Toco UA",
                    summary="оператор", content=f"- Пароль: {FAKE_PASSWORD_LINE}",
                    contradictions="", aliases=["ТОКО"], tags=[], sources=[],
                    status="quarantined")
    db.add(page)
    db.add(ChatLog(user_id=OWNER, role="bot", text=f"пароль {FAKE_PASSWORD_LINE}",
                   provider_eligible=False))
    await db.commit()

    report = await security_scan.run_scan(db, user_id=OWNER)
    assert report.documents_released == 1      # the password doc
    assert report.wiki_released == 1
    assert report.chat_released == 1
    assert report.documents_quarantined == 0   # token doc was already quarantined

    await db.refresh(pw_doc); await db.refresh(tok_doc); await db.refresh(page)
    assert pw_doc.status == "indexed"          # released, searchable again
    assert "security" not in (pw_doc.meta or {})
    assert tok_doc.status == "quarantined"     # token stays contained
    assert page.status == "active"

    # the password page is found again; its finding is resolved
    assert await wiki.find_page(db, OWNER, "personal", "ТОКО") is not None
    resolved = (await db.execute(select(SecurityFinding).where(
        SecurityFinding.resource_type == "document",
        SecurityFinding.resource_id == str(pw_doc.id)))).scalar_one()
    assert resolved.status == "resolved"
    # the token document keeps an OPEN finding
    open_tok = (await db.execute(select(SecurityFinding).where(
        SecurityFinding.resource_type == "document",
        SecurityFinding.resource_id == str(tok_doc.id)))).scalar_one()
    assert open_tok.status == "open"


@pytest.mark.asyncio
async def test_43_released_password_document_is_retrievable(db):
    """End to end: a password table quarantined by the old policy answers
    «який пароль до …» again after the reconciling scan."""
    from app.core.embeddings import get_embedder
    doc = Document(user_id=OWNER, domain="personal", title="Доступи операторів",
                   source_type="drive", source_ref="acc", content_hash="h-acc",
                   status="quarantined", chunk_count=1,
                   meta={"security": {"categories": ["password"],
                                      "finding_count": 1, "scanner_version": 1}})
    db.add(doc)
    await db.flush()
    text = "Toco UA (ТОКО) кабінет: логін i.k@example.invalid, пароль Qw3rty-Zx9-Lm"
    emb = (await get_embedder().embed([text]))[0]
    db.add(KnowledgeChunk(document_id=doc.id, user_id=OWNER, seq=0,
                          text=text, embedding=emb))
    await db.commit()

    assert await rag.retrieve(db, user_id=OWNER, domain="personal",
                              query="пароль Toco кабінет") == []
    await security_scan.run_scan(db, user_id=OWNER)
    hits = await rag.retrieve(db, user_id=OWNER, domain="personal",
                              query="пароль Toco кабінет")
    assert any("Qw3rty-Zx9-Lm" in h.text for h in hits)
